from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/caesar/Documents/ChatGPT/sackey")
OUT = ROOT / "deliverables"
SOURCE_CH12 = Path("/Users/caesar/Downloads/Bless_Kofi_Atisu_Chapters_One_and_Two.docx")
SOURCE_CH3 = Path("/Users/caesar/Downloads/Bless_Kofi_Atisu_Chapter_Three_Draft.docx")
OUT_CH12 = OUT / "Bless_Kofi_Atisu_Chapters_One_and_Two_Corrected.docx"
OUT_CH3 = OUT / "Bless_Kofi_Atisu_Chapter_Three_Design_Freeze_Draft.docx"
OUT_SPEC = OUT / "Bless_Kofi_Atisu_Implementation_Design_Specification.docx"

NAVY = "17365D"
BLUE = "D9E8F5"
ORANGE = "FBE5D6"
GREEN = "E2F0D9"
GREY = "E7E6E6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def style_cell(cell, *, bold=False, color="000000", size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.name = "Times New Roman"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)


def style_table(table, header_fill=NAVY, body_size=8.5) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    if table.rows:
        set_repeat_table_header(table.rows[0])
        set_row_cant_split(table.rows[0])
        for cell in table.rows[0].cells:
            set_cell_shading(cell, header_fill)
            style_cell(cell, bold=True, color="FFFFFF", size=body_size, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, row in enumerate(table.rows[1:], start=1):
        set_row_cant_split(row)
        for cell in row.cells:
            if i % 2 == 0:
                set_cell_shading(cell, "F3F6F8")
            style_cell(cell, size=body_size)


def set_run_font(run, size=12, bold=None, italic=None, color=None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_text(paragraph, text: str, size=12, bold=None, italic=None) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)


def replace_all_runs(doc: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                count += run.text.count(old)
                run.text = run.text.replace(old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text:
                            count += run.text.count(old)
                            run.text = run.text.replace(old, new)
    return count


def find_paragraph(doc: Document, startswith: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise KeyError(startswith)


def set_paragraph_by_start(
    doc: Document,
    startswith: str,
    text: str,
    *,
    size=12,
    bold=None,
    italic=None,
) -> None:
    p = find_paragraph(doc, startswith)
    set_paragraph_text(p, text, size=size, bold=bold, italic=italic)


def add_alt_text(doc: Document, descriptions: list[str]) -> None:
    drawings = doc.element.body.xpath(".//wp:docPr")
    for idx, node in enumerate(drawings):
        if idx < len(descriptions):
            node.set("descr", descriptions[idx])
            node.set("title", descriptions[idx].split(":", 1)[0])


def replace_embedded_image(doc: Document, image_index: int, image_path: Path) -> None:
    drawings = doc.element.body.xpath(".//w:drawing")
    drawing = drawings[image_index]
    blips = drawing.xpath(".//a:blip")
    if not blips:
        raise RuntimeError("No image relationship found")
    rid = blips[0].get(qn("r:embed"))
    image_part = doc.part.related_parts[rid]
    image_part._blob = image_path.read_bytes()


def svg_to_png(svg_path: Path, png_path: Path) -> None:
    # Render this precise box-and-arrow academic diagram with Pillow so labels
    # remain deterministic and are not regenerated by a probabilistic image tool.
    from PIL import Image, ImageDraw, ImageFont

    scale = 1.5
    width, height = int(1200 * scale), int(820 * scale)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"

    def font(size, bold=False):
        path = bold_path if bold else font_path
        return ImageFont.truetype(path, int(size * scale))

    def xy(values):
        return tuple(int(v * scale) for v in values)

    def centred(text, center_x, y, size=14, bold=False, fill="#333333"):
        f = font(size, bold)
        box = draw.textbbox((0, 0), text, font=f)
        x = int(center_x * scale - (box[2] - box[0]) / 2)
        draw.text((x, int(y * scale)), text, font=f, fill=fill)

    def box(coords, fill, outline="#17365D", radius=14, width=3):
        draw.rounded_rectangle(xy(coords), radius=int(radius * scale), fill=fill, outline=outline, width=int(width * scale))

    def arrow(start, end, fill="#17365D"):
        draw.line(xy((*start, *end)), fill=fill, width=int(3 * scale))
        x2, y2 = end
        draw.polygon([xy((x2, y2)), xy((x2 - 8, y2 - 11)), xy((x2 + 8, y2 - 11))], fill=fill)

    centred("WELCOME AND ACCESS LAYER", 600, 20, 24, True, "#17365D")
    box((110, 55, 1090, 147), "#D9E8F5")
    centred("ATU welcome interface and General Help", 600, 73, 17, True, "#17365D")
    centred("Student/Lecturer registration, sign-in, password recovery and pending approval", 600, 101)
    centred("Role verification and academic setup", 600, 126)
    draw.line(xy((600, 147, 600, 184)), fill="#17365D", width=int(3 * scale))
    draw.line(xy((250, 184, 950, 184)), fill="#17365D", width=int(3 * scale))
    for x in (250, 600, 950):
        arrow((x, 184), (x, 210))

    role_boxes = [
        (80, 216, 420, 454, "#FBE5D6", "ADMINISTRATOR", ["Accounts and approvals", "Academic structure and access", "Approved-resource catalogue", "Audit, reports and notifications", "Secure password resets"]),
        (430, 216, 770, 454, "#D9E8F5", "LECTURER", ["Practical task creation", "Approved-resource selection", "Submission and evidence review", "Grading, feedback and release", "Notifications and exports"]),
        (780, 216, 1120, 454, "#E2F0D9", "STUDENT", ["Assigned practical tasks", "Approved demonstration resources", "Saved sessions and evidence", "Submission and resubmission", "Results, notifications and profile"]),
    ]
    for x1, y1, x2, y2, fill, title, lines in role_boxes:
        box((x1, y1, x2, y2), fill)
        cx = (x1 + x2) / 2
        centred(title, cx, 235, 17, True, "#17365D")
        for idx, line in enumerate(lines):
            centred(line, cx, 279 + idx * 25)
        arrow((cx, 454), (cx, 493), "#E87524")

    box((95, 500, 1105, 605), "#FFF2CC", "#E87524")
    centred("VIRTUAL MODELLING AND DEMONSTRATION PAGE", 600, 520, 17, True, "#C65911")
    centred("Approved technical demonstrations, simulations and defined practical resources", 600, 554)
    centred("Run ID, resource/version, parameters, plots, metrics, files, animation evidence and submission status", 600, 580, 12)
    arrow((600, 605), (600, 644))
    box((120, 650, 1080, 762), "#E7E6E6", "#666666")
    centred("SHARED DATA AND RECORD LAYER", 600, 673, 17, True)
    centred("Users and roles | Faculties, departments and courses | Practical questions and assignments", 600, 708)
    centred("Submissions and grades | Practical files and outputs | Notifications and audit records", 600, 736)
    image.save(png_path, format="PNG", optimize=True)


def build_architecture_figure() -> Path:
    assets = OUT / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    svg_path = assets / "figure_1_1_corrected.svg"
    png_path = assets / "figure_1_1_corrected.png"
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="820" viewBox="0 0 1200 820">
<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#17365D"/></marker></defs>
<rect width="1200" height="820" fill="white"/>
<style>.title{{font:700 24px Arial;fill:#17365D}} .h{{font:700 17px Arial}} .t{{font:14px Arial;fill:#333}} .small{{font:12px Arial;fill:#444}}</style>
<text x="600" y="38" text-anchor="middle" class="title">WELCOME AND ACCESS LAYER</text>
<rect x="110" y="55" width="980" height="92" rx="15" fill="#D9E8F5" stroke="#17365D" stroke-width="3"/>
<text x="600" y="84" text-anchor="middle" class="h" fill="#17365D">ATU welcome interface and General Help</text>
<text x="600" y="110" text-anchor="middle" class="t">Student/Lecturer registration, sign-in, password recovery and pending approval</text>
<text x="600" y="133" text-anchor="middle" class="t">Role verification and academic setup</text>
<line x1="600" y1="147" x2="600" y2="184" stroke="#17365D" stroke-width="3" marker-end="url(#arrow)"/>
<line x1="250" y1="184" x2="950" y2="184" stroke="#17365D" stroke-width="3"/>
<line x1="250" y1="184" x2="250" y2="210" stroke="#17365D" stroke-width="3" marker-end="url(#arrow)"/>
<line x1="600" y1="184" x2="600" y2="210" stroke="#17365D" stroke-width="3" marker-end="url(#arrow)"/>
<line x1="950" y1="184" x2="950" y2="210" stroke="#17365D" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="80" y="216" width="340" height="238" rx="14" fill="#FBE5D6" stroke="#17365D" stroke-width="3"/>
<text x="250" y="250" text-anchor="middle" class="h" fill="#17365D">ADMINISTRATOR</text>
<text x="250" y="284" text-anchor="middle" class="t">Accounts and approvals</text><text x="250" y="309" text-anchor="middle" class="t">Academic structure and access</text><text x="250" y="334" text-anchor="middle" class="t">Approved-resource catalogue</text><text x="250" y="359" text-anchor="middle" class="t">Audit, reports and notifications</text><text x="250" y="384" text-anchor="middle" class="t">Secure password resets</text>
<rect x="430" y="216" width="340" height="238" rx="14" fill="#D9E8F5" stroke="#17365D" stroke-width="3"/>
<text x="600" y="250" text-anchor="middle" class="h" fill="#17365D">LECTURER</text>
<text x="600" y="284" text-anchor="middle" class="t">Practical task creation</text><text x="600" y="309" text-anchor="middle" class="t">Approved-resource selection</text><text x="600" y="334" text-anchor="middle" class="t">Submission and evidence review</text><text x="600" y="359" text-anchor="middle" class="t">Grading, feedback and release</text><text x="600" y="384" text-anchor="middle" class="t">Notifications and exports</text>
<rect x="780" y="216" width="340" height="238" rx="14" fill="#E2F0D9" stroke="#17365D" stroke-width="3"/>
<text x="950" y="250" text-anchor="middle" class="h" fill="#17365D">STUDENT</text>
<text x="950" y="284" text-anchor="middle" class="t">Assigned practical tasks</text><text x="950" y="309" text-anchor="middle" class="t">Approved demonstration resources</text><text x="950" y="334" text-anchor="middle" class="t">Saved sessions and evidence</text><text x="950" y="359" text-anchor="middle" class="t">Submission and resubmission</text><text x="950" y="384" text-anchor="middle" class="t">Results, notifications and profile</text>
<line x1="250" y1="454" x2="250" y2="493" stroke="#E87524" stroke-width="3" marker-end="url(#arrow)"/><line x1="600" y1="454" x2="600" y2="493" stroke="#E87524" stroke-width="3" marker-end="url(#arrow)"/><line x1="950" y1="454" x2="950" y2="493" stroke="#E87524" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="95" y="500" width="1010" height="105" rx="15" fill="#FFF2CC" stroke="#E87524" stroke-width="3"/>
<text x="600" y="535" text-anchor="middle" class="h" fill="#C65911">VIRTUAL MODELLING AND DEMONSTRATION PAGE</text>
<text x="600" y="566" text-anchor="middle" class="t">Approved technical demonstrations, simulations and defined practical resources</text>
<text x="600" y="589" text-anchor="middle" class="small">Run ID, resource/version, parameters, plots, metrics, files, animation evidence and submission status</text>
<line x1="600" y1="605" x2="600" y2="644" stroke="#17365D" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="120" y="650" width="960" height="112" rx="15" fill="#E7E6E6" stroke="#666" stroke-width="3"/>
<text x="600" y="686" text-anchor="middle" class="h" fill="#333">SHARED DATA AND RECORD LAYER</text>
<text x="600" y="718" text-anchor="middle" class="t">Users and roles | Faculties, departments and courses | Practical questions and assignments</text>
<text x="600" y="743" text-anchor="middle" class="t">Submissions and grades | Practical files and outputs | Notifications and audit records</text>
</svg>'''
    svg_path.write_text(svg, encoding="utf-8")
    svg_to_png(svg_path, png_path)
    return png_path


def correct_chapters_one_two() -> None:
    doc = Document(SOURCE_CH12)

    replacements = {
        "approved Practical Research Models": "approved demonstration resources",
        "Automobile/Automotive Engineering": "Automotive Engineering",
        "automobile/automotive engineering": "automotive engineering",
        "simulation files and practical-output data": "approved demonstration files and practical-output data",
        "demonstrations, simulations and other practical concepts": "accessing approved simulations, demonstrations or practical concepts",
        "which models require improvement": "which approved resources may require review by their responsible providers",
        "practical-model integration": "approved-resource integration",
        "practical models": "approved demonstration resources",
        "integrated models": "integrated resources",
        "model interfaces": "resource interfaces",
        "continued model development": "continued resource development",
        "Labster-style simulations": "gamified laboratory simulations",
        "World Wide Web Consortium, 2023": "World Wide Web Consortium, 2024",
        "World Wide Web Consortium. (2023).": "World Wide Web Consortium. (2024).",
    }
    for old, new in replacements.items():
        replace_all_runs(doc, old, new)

    set_paragraph_by_start(
        doc,
        "The aim of this study",
        "The aim of this study is to develop a web-based platform for managing automotive practical assessment, including question setting, student submission, result organisation, and linkage to approved simulation or practical-task outputs.",
    )
    set_paragraph_by_start(
        doc,
        "To provide an integrated platform for accessing approved simulations",
        "To provide an integrated platform for accessing approved simulations, demonstrations or practical concepts that can support automotive practical learning.",
    )
    set_paragraph_by_start(
        doc,
        "How can a database system be developed",
        "How can a database system be developed to store and organise practical questions, student responses, assessment records, approved demonstration files and practical-output data?",
    )
    set_paragraph_by_start(
        doc,
        "How can accessing approved simulations",
        "How can approved simulations, demonstrations or practical concepts be accessed through an integrated platform to support automotive practical learning?",
    )
    set_paragraph_by_start(
        doc,
        "Automotive education is useful only when",
        "Automotive education fulfils its practical purpose when knowledge can be carried into purposeful action. A student may be able to describe a braking system, an electrical charging circuit or an engine-management principle and still struggle to diagnose a fault, select a safe procedure or explain the meaning of an observed result. This is why engineering laboratory education has traditionally been treated as more than illustration. Feisel and Rosa (2005) identify objectives that include instrumentation, experimentation, data analysis, design, learning from failure, creativity and psychomotor development. These objectives show that competence depends on what the learner can do with knowledge, not only on what the learner can recall.",
    )

    p = find_paragraph(doc, "The third limitation is that the educational strength")
    if "not certified" not in p.text:
        new_text = p.text + " The connected demonstrations are educational tools; they are not certified industrial-design, homologation or validation systems and must not be used as the sole basis for safety-critical decisions."
        set_paragraph_text(p, new_text)

    set_paragraph_by_start(
        doc,
        "The gap is therefore an integration and management gap",
        "The gap is therefore an integration and management gap, not a claim that virtual laboratories or learning systems do not exist. Among the systems reviewed, there is limited evidence of a focused web-based workflow that treats the practical task as the central academic record; distinguishes Administrator, Lecturer and Student responsibilities; connects only approved demonstration resources; captures a reviewable work trail; supports formative assessment; and remains usable, accessible and secure. The proposed study addresses this gap within automotive practical education.",
    )
    set_paragraph_by_start(
        doc,
        "The study is organised into five chapters.",
        "The study is organised into five chapters. Chapter One presents the background, problem statement, aim, specific objectives, research questions, significance, scope and limitations. Chapter Two reviews the related literature and develops the conceptual basis of the study. Chapter Three presents the system design and planning, proposed development and implementation procedure, and testing and evaluation method. Chapter Four presents and evaluates the implemented platform, while Chapter Five summarises the findings, conclusions and recommendations.",
    )
    set_paragraph_by_start(
        doc,
        "The identified gap is the absence of a focused",
        "The identified gap is the limited evidence, among the reviewed systems, of a focused and integrated workflow for managing automotive practical education from approved access through task completion, evidence, feedback and oversight. The proposed platform responds to this gap without claiming to replace the physical workshop or reproduce the internal operation of every demonstration tool. The next chapter will present the system design and planning, proposed development and implementation procedure, and testing and evaluation method.",
    )

    figure_path = build_architecture_figure()
    replace_embedded_image(doc, 0, figure_path)
    add_alt_text(
        doc,
        [
            "Figure 1.1: Conceptual architecture showing the access layer, Administrator, Lecturer and Student workspaces, the Virtual Modelling and Demonstration page, and the shared record layer.",
            "Figure 2.1: Conceptual framework linking inputs through web-based management to educational outputs, supported by usability, accessibility, security and resource quality.",
        ],
    )

    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_CH12)


def rewrite_chapter_three() -> None:
    doc = Document(SOURCE_CH3)

    doc.tables[0].cell(1, 3).text = (
        "After Administrator setup of the course offering and resource access, the Lecturer completes the create-to-release workflow independently"
    )
    doc.tables[1].cell(1, 2).text = (
        "Cannot alter a Lecturer’s academic judgement; any authorised correction follows a documented academic process and remains auditable"
    )
    style_table(doc.tables[0], body_size=7.8)
    style_table(doc.tables[1], body_size=8.3)

    set_paragraph_by_start(
        doc,
        "The architecture also supports later hosting.",
        "The architecture also supports later hosting. During development, the Django application and PostgreSQL database will run locally through a documented environment. For deployment, the application service, database, protected media storage and web proxy can be placed on an institutional server or approved cloud host. The eventual hosting location may change, but the logical relationship among the layers will remain the same.",
    )

    set_paragraph_by_start(
        doc,
        "Development translates the approved design",
        "After approval of the design baseline, development will translate the requirements into a working web application. A Django modular monolith will contain the server-rendered interface, application rules, integration endpoints and data-access layer within one controlled codebase, while PostgreSQL will store relational records and protected media storage will hold uploaded evidence. Development will proceed in small vertical increments. The first increment will complete registration, validation, pending status, Administrator approval, academic setup and the sign-in restriction before practical-task functions are added.",
    )
    set_paragraph_by_start(
        doc,
        "The implementation basis is open source",
        "The proposed implementation basis is open source and can run on an ordinary development computer. The selected stack is shown in Table 3.4. Python 3.13 and Django 5.2 LTS will provide the application framework, server-rendered interface, authentication and permission foundation. Django templates, HTMX 2.x and Bootstrap 5.3 will support responsive interaction without requiring a separate browser application. Django REST Framework 3.16 will be limited to documented integration endpoints, while PostgreSQL 18 will preserve the relational records. Exact dependency versions will be locked in the project environment and container image tags so that a later installation does not silently use incompatible releases (Django Software Foundation, 2025a, 2025b; Django REST Framework, 2025; PostgreSQL Global Development Group, 2026).",
    )

    table = doc.tables[3]
    rows = [
        ("Application framework", "Python 3.13 with Django 5.2 LTS", "Server-rendered web application, authentication, permissions, validation and application services"),
        ("Interface", "HTML5, Django templates, HTMX 2.x and Bootstrap 5.3", "Responsive Administrator, Lecturer and Student workspaces with focused partial-page interaction"),
        ("Integration API", "Django REST Framework 3.16", "Documented endpoints only for approved-resource launch and agreed result metadata"),
        ("Database", "PostgreSQL 18 with Django migrations", "Relational storage, constraints, transactions and indexed retrieval"),
        ("File handling", "Private media directory in development; S3-compatible option for deployment", "Stores evidence outside the public static path while the database keeps metadata and ownership"),
        ("Deployment", "Docker Compose; Gunicorn and Nginx for hosted deployment", "Repeatable local services and a production-ready application boundary"),
        ("Source control", "Git and GitHub", "Version history, tagged study release and reproducible source repository"),
        ("Testing", "pytest, pytest-django, Playwright and axe-core", "Unit, integration, end-to-end and automated accessibility checks"),
        ("Design and inspection", "diagrams.net and browser developer tools", "Architecture diagrams, responsive inspection and browser-level diagnostics"),
    ]
    for i, (a, b, c) in enumerate(rows, start=1):
        table.cell(i, 0).text = a
        table.cell(i, 1).text = b
        table.cell(i, 2).text = c
    style_table(table)

    set_paragraph_by_start(doc, "3.3.1 Source Organisation", "3.3.1 Source Organisation and Interface Implementation Plan", bold=True)
    set_paragraph_by_start(
        doc,
        "The source repository is organised",
        "The source repository will be organised as a Django project with configuration, domain applications, templates, static assets, protected-media handling, tests and documentation separated by responsibility. Domain applications will cover accounts, academics, resources, assessments, submissions, results, notifications, reports and audit records. Database migrations and controlled starter data will be retained with the source so that the schema can be recreated in a known order. Tests will mirror the major workflows, while the documentation area will contain environment, deployment and recovery instructions.",
    )
    set_paragraph_by_start(
        doc,
        "The frontend is implemented as role-aware routes.",
        "The interface will use role-aware Django views and templates. The authenticated session will contain only the information needed to identify the user and enforce access; the server will remain the authority for every protected request. Reusable template components will be used for form fields, status labels, notification items, tables, confirmation prompts and empty states. HTMX will be applied only where partial-page interaction reduces unnecessary navigation, and every operation will retain a normal server-side fallback.",
    )
    set_paragraph_by_start(
        doc,
        "Responsive behaviour is achieved",
        "Responsive behaviour will be achieved through flexible layout rules and breakpoints rather than separate applications for wide and narrow screens. The navigation panel will remain visible on wider screens and become a controlled drawer on narrow screens. Tables containing academic records will retain essential columns and allow horizontal movement when necessary, while forms will use a single-column arrangement on small screens. Error messages will appear close to the affected field and will also be summarised at the beginning of longer forms.",
    )

    set_paragraph_by_start(doc, "3.3.2 Application Services", "3.3.2 Application Services and Database Implementation Plan", bold=True)
    set_paragraph_by_start(
        doc,
        "The Express server exposes",
        "Django views and application services will handle authentication, users, academic setup, resources, tasks, practical sessions, submissions, grades, notifications, reports and audit records. Each protected request will pass through session authentication, server-side permission checks, input validation and domain rules before a state change is committed. Django REST Framework endpoints will be exposed only where an approved external resource requires a documented exchange. Multi-record operations will use database transactions so that related records either succeed together or leave no partial academic state.",
    )
    set_paragraph_by_start(
        doc,
        "Versioned SQL migration files create",
        "Versioned Django migration files will create the database in a known sequence. Primary keys will uniquely identify records; foreign keys will prevent orphaned submissions or grades; unique constraints will prevent duplicate enrolment; and indexes will support common searches such as open tasks by course, submissions awaiting review and audit events by date. A controlled management command will create the initial academic structure where required. The bootstrap Administrator account will be created through the deployment process, never through public registration, and its temporary credential must be changed at first sign-in.",
    )
    set_paragraph_by_start(
        doc,
        "Uploaded evidence is accepted",
        "Uploaded evidence will be accepted only after the server confirms the user, task, deadline and permitted file rules. The stored file will receive a generated name, while the original name will be retained as metadata for display. Type, size and ownership will be checked before download, and the file will never be exposed through a guessable public path. A checksum may be retained to support integrity checks and duplicate detection.",
    )

    set_paragraph_by_start(doc, "3.3.3 Implementation", "3.3.3 Planned Implementation of the Role-Based Modules", bold=True)
    set_paragraph_by_start(
        doc,
        "The Administrator module implements",
        "The Administrator module will implement registration approval, account suspension, password-reset support, academic structure, course access, approved-resource registration, notifications, reports and the activity log. Academic setup will follow a hierarchy so that a course offering cannot be created without a course and academic period. Resource registration will record the responsible provider, version, launch method, status, permitted data exchange and any assumptions that users must see.",
    )
    set_paragraph_by_start(
        doc,
        "The Lecturer module implements",
        "The Lecturer module will implement task creation and management. A Lecturer will select an assigned course, enter the practical question, state the expected output, attach a rubric or marking guide, choose an approved resource if required and define the deadline. Once published, the task will become visible to the relevant Students. The review area will group submissions by task and retain the Student’s saved session, uploaded evidence, explanations, feedback, revisions and final decision. Results may be released and exported without exposing another course’s records.",
    )
    set_paragraph_by_start(
        doc,
        "The Student module presents",
        "The Student module will present only tasks linked to the Student’s enrolments. The Student will be able to read the instructions, open the approved resource, save permitted session information, upload evidence and submit the response. Where resubmission is allowed, the new attempt will be created as another revision rather than replacing the first. Feedback and results will appear only when the Lecturer releases them. Notifications will draw attention to a new task, approaching deadline, returned feedback or released result.",
    )
    set_paragraph_by_start(
        doc,
        "Profile management and role-specific logout are shared functions.",
        "Profile management and role-specific logout will be shared functions. The proposed five-minute inactivity rule will be enforced at the protected-session level, so activity on an unprotected page will not indefinitely keep an academic session alive. On expiry, the user will be returned to sign-in and the attempted action will not be performed silently.",
    )

    set_paragraph_by_start(doc, "3.3.4 Integration", "3.3.4 Planned Integration of Approved Practical Resources", bold=True)
    p = find_paragraph(doc, "A practical resource is integrated")
    set_paragraph_text(p, p.text.replace("is integrated", "will be integrated", 1).replace("The Administrator records", "The Administrator will record", 1).replace("The resource is then available", "The resource will then be available", 1).replace("The lecturer links", "The Lecturer will link", 1))
    p = find_paragraph(doc, "For a simple resource, integration may")
    set_paragraph_text(p, p.text.replace("the platform can issue", "the platform may issue").replace("The returning resource supplies", "The returning resource will supply").replace("The platform validates", "The platform will validate"))

    set_paragraph_by_start(doc, "3.3.5 Security", "3.3.5 Security and Data-Protection Implementation Plan", bold=True)
    set_paragraph_by_start(
        doc,
        "Passwords are stored only",
        "Passwords will be stored only as salted hashes using Django’s password-hashing framework, with Argon2 configured as the preferred hasher. Authenticated sessions will use random server-managed identifiers in HttpOnly, Secure and SameSite cookies. Role permission will be checked for every protected view or API operation, and ownership or course membership will be checked again when a record is requested. This will prevent a user from gaining access simply by changing a page address or record identifier.",
    )
    p = find_paragraph(doc, "The server validates field length")
    set_paragraph_text(p, p.text.replace("The server validates", "The server will validate", 1).replace("Database operations use", "Database operations will use", 1).replace("browser output is encoded", "browser output will be encoded", 1).replace("File uploads use", "File uploads will use", 1).replace("attempts are rate-limited", "attempts will be rate-limited", 1).replace("reset tokens expire", "reset tokens will expire", 1).replace("secret values are loaded", "secret values will be loaded", 1))
    p = find_paragraph(doc, "Audit events are created")
    set_paragraph_text(p, p.text.replace("Audit events are created", "Audit events will be created", 1).replace("Routine page views are not treated", "Routine page views will not be treated", 1).replace("Logs avoid storing", "Logs will avoid storing", 1))

    set_paragraph_by_start(doc, "3.3.6 Reproducible", "3.3.6 Reproducible Setup and Deployment Plan", bold=True)
    set_paragraph_by_start(
        doc,
        "Docker Compose is used",
        "Docker Compose will describe the Django application and PostgreSQL services in one configuration. It is suited to this purpose because it defines the services, network, volumes and environment needed to start the application consistently (Docker, 2026). Gunicorn and Nginx will form the hosted web boundary, while the same repository may also be run through Django’s development server for local work. The container route will be the reference procedure for reproduction.",
    )
    set_paragraph_by_start(
        doc,
        "The reproduction package contains",
        "The reproduction package will contain the source code, locked Python dependencies, Dockerfiles, Compose configuration, Django migrations, controlled starter-data commands, an environment template, test commands and README instructions. Personal credentials, uploaded student evidence and production secrets will be excluded. A clean installation will follow the steps below.",
    )
    setup_steps = [
        "Install Git and Docker Desktop, including Docker Compose, and confirm that both commands are available.",
        "Clone the tagged project repository and remain on the documented release tag used for the study.",
        "Copy the environment template to the local environment file, create new secret values and keep production credentials outside the repository.",
        "Run docker compose up --build so that the Django application and PostgreSQL services are created from the declared images and dependencies.",
        "Run the versioned Django migrations and controlled starter-data command, then create the bootstrap Administrator account through the documented deployment command.",
        "Open the local application address, normally http://localhost:8000, and confirm that the application and database health checks report a ready state.",
        "Run the pytest and Playwright commands and retain the generated reports before adding real users or practical resources.",
        "Create a separate backup volume or database dump before significant demonstration data is entered, and record any local configuration change in the README.",
    ]
    for old_start, new_text in zip(
        ["Install Git", "Clone the tagged", "Copy the environment", "Run docker compose", "Run the versioned", "Open the local", "Run the automated", "Create a separate backup"],
        setup_steps,
    ):
        set_paragraph_by_start(doc, old_start, new_text)

    set_paragraph_by_start(
        doc,
        "After the containers start",
        "After the containers start, the operator will sign in with the bootstrap Administrator account, change the temporary password, confirm the academic structure and create or approve the first Lecturer and Student accounts. The same setup may remain on localhost for demonstration. Later institutional hosting will require HTTPS, a managed domain, regular database backup, monitored storage and a documented recovery procedure, but will not require a change to the logical architecture.",
    )

    set_paragraph_by_start(
        doc,
        "Testing checks whether",
        "Testing will check whether the implemented platform behaves as the approved design requires. It will proceed from small code units to the complete user workflow. Unit tests will check validation rules and permission decisions. Django integration tests will check that related records are created or rejected correctly. Playwright end-to-end tests will operate the interface as an Administrator, Lecturer or Student. Manual evaluation will then examine clarity, accessibility and usability in tasks that automated tests cannot judge completely.",
    )
    set_paragraph_by_start(
        doc,
        "Testing is performed on a clean local deployment",
        "Testing will be performed on a clean local deployment created from the same reproduction instructions. Test data will use fictional users and non-sensitive files. Each case will record the requirement, starting state, actions, expected result, observed result, evidence and correction where necessary. Actual outcomes and screenshots will be presented in Chapter Four; the present chapter defines the procedure and acceptance basis.",
    )
    replace_all_runs(doc, "Automated axe-core scans", "Automated axe-core scans")
    replace_all_runs(doc, "desktop and mobile-width layouts", "wide-screen and narrow-screen layouts")
    table = doc.tables[4]
    table.cell(8, 3).text = "Manual checklist and axe-core report"
    table.cell(9, 1).text = "Repeat core workflow in current Chrome, Edge and Firefox at wide and narrow screen widths"
    style_table(table)

    set_paragraph_by_start(
        doc,
        "This chapter has presented a focused method",
        "This chapter has presented a focused method for the design, proposed development and evaluation of the web-based Virtual Laboratory Management Platform. The design begins with the approved objectives and assigns clear responsibilities to the Administrator, Lecturer and Student. A three-layer logical architecture separates the interface, application rules and records, while approved practical resources remain connected but technically independent. The relational data model preserves the route from academic setup and task publication through practical evidence, feedback and result release.",
    )
    set_paragraph_by_start(
        doc,
        "Development is based on a reproducible React",
        "Development will use a reproducible Django 5.2 LTS and PostgreSQL 18 modular-monolith stack supported by Git, locked dependencies and Docker Compose. The planned implementation covers role-based workspaces, protected resource integration, structured evidence, assessment, notifications, session control and audit records. The reproduction procedure identifies the files and steps required to rebuild the system without relying on the original development computer.",
    )
    set_paragraph_by_start(
        doc,
        "Testing follows the same objectives",
        "Testing will follow the same objectives and workflows. This chapter defines the functional, integration, security, accessibility, compatibility, performance and usability checks. Chapter Four will present only the interfaces actually implemented, measured results, test evidence and recorded corrections.",
    )

    # Rebuild the reference list so obsolete React/Node/Express/Vite sources are removed.
    heading = find_paragraph(doc, "CHAPTER THREE REFERENCES")
    heading.paragraph_format.page_break_before = True
    body = doc._body._element
    seen = False
    for child in list(body):
        if child is heading._element:
            seen = True
            continue
        if seen and child.tag == qn("w:p"):
            body.remove(child)
    references = [
        "Brooke, J. (1996). SUS: A ‘quick and dirty’ usability scale. In P. W. Jordan, B. Thomas, B. A. Weerdmeester, & I. L. McClelland (Eds.), Usability evaluation in industry (pp. 189–194). Taylor & Francis.",
        "Django REST Framework. (2025). Version 3.16 announcement. https://www.django-rest-framework.org/community/3.16-announcement/",
        "Django Software Foundation. (2025a). Django 5.2 release notes. https://docs.djangoproject.com/en/5.2/releases/5.2/",
        "Django Software Foundation. (2025b). Using the Django authentication system. https://docs.djangoproject.com/en/5.2/topics/auth/default/",
        "Docker. (2026). Docker Compose. https://docs.docker.com/compose/",
        "International Organization for Standardization. (2018). ISO 9241-11:2018: Ergonomics of human-system interaction—Part 11: Usability: Definitions and concepts. ISO.",
        "OWASP Foundation. (2025a). OWASP Application Security Verification Standard 5.0.0. https://owasp.org/www-project-application-security-verification-standard/",
        "OWASP Foundation. (2025b). OWASP Top 10:2025. https://owasp.org/Top10/2025/",
        "Peffers, K., Tuunanen, T., Rothenberger, M. A., & Chatterjee, S. (2007). A design science research methodology for information systems research. Journal of Management Information Systems, 24(3), 45–77. https://doi.org/10.2753/MIS0742-1222240302",
        "PostgreSQL Global Development Group. (2026). PostgreSQL versioning policy. https://www.postgresql.org/support/versioning/",
        "Souppaya, M., Scarfone, K., & Dodson, D. (2022). Secure Software Development Framework (SSDF) version 1.1: Recommendations for mitigating the risk of software vulnerabilities (NIST Special Publication 800-218). National Institute of Standards and Technology. https://doi.org/10.6028/NIST.SP.800-218",
        "World Wide Web Consortium. (2024). Web Content Accessibility Guidelines (WCAG) 2.2. https://www.w3.org/TR/WCAG22/",
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        set_paragraph_text(p, ref, size=10.5)

    add_alt_text(
        doc,
        [
            "Figure 3.1: Iterative cycle connecting design and planning, development and implementation, and testing and evaluation.",
            "Figure 3.2: Six-stage practical-assessment workflow from administrative approval through preserved results and audit records.",
            "Figure 3.3: Three-layer architecture with the approved external practical-resource boundary.",
            "Figure 3.4: Logical data-model groups connecting users, academics, tasks, resources, evidence, assessment and oversight.",
        ],
    )

    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_CH3)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=10)


def add_bullet(doc: Document, text: str, level=0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)


def add_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]], widths=None, size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    style_table(table, body_size=size)
    return table


def add_heading(doc: Document, text: str, level=1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)


def build_specification() -> None:
    doc = Document()
    configure_document(doc)
    add_page_number(doc.sections[0].footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run("ATU VIRTUAL LABORATORY MANAGEMENT PLATFORM")
    set_run_font(r, size=18, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("IMPLEMENTATION DESIGN SPECIFICATION")
    set_run_font(r, size=16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(32)
    r = p.add_run("Student: Bless Kofi Atisu\nProgramme: MTech Automotive Engineering\nInstitution: Accra Technical University")
    set_run_font(r, size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    r = p.add_run("Design-freeze candidate v1.0 — September 2026")
    set_run_font(r, size=11, italic=True, color="666666")
    doc.add_page_break()

    add_heading(doc, "1. Purpose and Status", 1)
    add_para(doc, "This specification converts the approved study aim and four objectives into a buildable system baseline. It defines the platform boundary, roles, data relationships, workflow states, pages, security controls, integration contract and acceptance tests. It is a design-freeze candidate: implementation should begin only after the decisions in Section 14 are accepted. The document records intended behaviour and must not be treated as evidence that the platform or its test results already exist.")

    add_heading(doc, "2. Approved Study Boundary", 1)
    add_para(doc, "The product is a web-based educational management platform for automotive practical assessment. It manages questions, approved resource access, practical sessions, student evidence, submission and resubmission, lecturer assessment, results, notifications, reports and audit records. Approved simulations or demonstrations remain technically independent resources. The platform may launch them and receive agreed metadata, but it does not create or validate their internal engineering equations.")
    add_heading(doc, "2.1 Included", 2)
    for item in [
        "Public Student and Lecturer registration followed by pending approval and academic setup.",
        "Administrator-only account creation through deployment or an authorised administrative process.",
        "Faculties, departments, programmes, courses, academic periods, course offerings, teaching assignments and enrolments.",
        "Practical tasks, questions, learning outcomes, deadlines, rubrics, evidence requirements and approved-resource links.",
        "Practical sessions, run identifiers, resource versions, selected parameters, status, outputs, files and submission revisions.",
        "Lecturer feedback, grades, result release, exports, notifications and attributable audit events.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2.2 Excluded", 2)
    for item in [
        "Industrial research management, calibration, homologation, certification or safety-critical control.",
        "Scientific creation or validation of every external model, simulation or demonstration.",
        "Replacement of ATU workshops, physical laboratories, supervised tool handling or on-vehicle diagnosis.",
        "Public Administrator registration and unauthorised alteration of a Lecturer’s academic judgement.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Technology and Architecture Decision", 1)
    add_table(
        doc,
        ["Layer", "Selected technology", "Design decision"],
        [
            ("Application", "Python 3.13; Django 5.2 LTS", "Modular monolith containing server-rendered pages, business rules, authentication, permissions and data access."),
            ("Interaction", "Django templates; HTMX 2.x; Bootstrap 5.3", "Responsive web interface with progressive enhancement and normal server-side fallbacks."),
            ("Integration", "Django REST Framework 3.16", "Used only for documented approved-resource launch and result-return endpoints."),
            ("Data", "PostgreSQL 18; Django migrations", "Relational constraints, transactions, indexed queries and reproducible schema changes."),
            ("Files", "Protected local media in development; S3-compatible deployment option", "Evidence is never served from a guessable public static path."),
            ("Hosting", "Docker Compose; Gunicorn; Nginx", "Repeatable development services and a clear hosted boundary with HTTPS termination."),
            ("Verification", "pytest; pytest-django; Playwright; axe-core", "Unit, integration, end-to-end, permissions and accessibility screening."),
        ],
        widths=[1.1, 2.0, 3.6],
    )
    add_para(doc, "Architecture rule: the browser never decides permission. Every protected view, object retrieval, file download and state-changing request must be authorised on the server against the user’s role, course relationship and record ownership.")

    add_heading(doc, "4. Roles and Permission Matrix", 1)
    permission_rows = [
        ("Public visitor", "View Welcome and General Help; register as Student or Lecturer; sign in; request password recovery", "No academic records or protected resources"),
        ("Administrator", "Approve/suspend users; manage academic hierarchy and course access; register/approve resources; support password resets; review reports, notifications and audit records", "Cannot register publicly or alter Lecturer academic judgement"),
        ("Lecturer", "Manage tasks for assigned course offerings; select approved resources; define evidence; review submissions; give feedback; grade; release and export results", "No system-wide user administration; no access to unassigned courses"),
        ("Student", "View enrolled tasks; open approved resources; save sessions and evidence; submit/resubmit; read released feedback and results; manage profile", "No task creation, grading, academic configuration or access to another Student’s work"),
    ]
    add_table(doc, ["Role", "Permitted actions", "Boundary"], permission_rows, widths=[1.15, 3.9, 1.8], size=8.2)

    add_heading(doc, "5. Account and Academic Workflows", 1)
    add_heading(doc, "5.1 Account lifecycle", 2)
    add_table(
        doc,
        ["Current state", "Allowed transition", "Trigger and rule"],
        [
            ("Pending approval", "Rejected or Academic setup required", "Administrator verifies identity and selected role; sign-in remains blocked until approval."),
            ("Academic setup required", "Active", "Administrator assigns the approved faculty/department/programme and relevant course access."),
            ("Active", "Suspended", "Administrator records a reason; protected access stops while historical records remain."),
            ("Suspended", "Active", "Authorised reinstatement creates a new audit event."),
            ("Rejected", "Pending approval", "A corrected registration may be reconsidered without becoming an Administrator account."),
        ],
        widths=[1.5, 1.8, 3.6],
    )
    add_heading(doc, "5.2 Academic hierarchy", 2)
    add_para(doc, "The permanent hierarchy is Faculty → Department → Programme → Course. Delivery is represented separately as Academic Period → Course Offering. A Lecturer receives a Teaching Assignment to a Course Offering, while a Student receives an Enrolment. This separation preserves historical records when curricula, staff or semesters change.")
    add_heading(doc, "5.3 Principal academic sequence", 2)
    for step in [
        "Administrator approves the Lecturer and Student and completes academic access.",
        "Administrator registers and approves an external demonstration resource.",
        "Lecturer creates a draft task for an assigned course offering, sets the deadline and evidence requirements, and links an approved resource where needed.",
        "Lecturer publishes the task; enrolled Students receive access and a notification.",
        "Student opens the task, starts or resumes a practical session, records permitted parameters and outputs, uploads evidence and confirms submission.",
        "Lecturer reviews the preserved work trail, returns feedback or records a grade, and releases the result when ready.",
        "Notifications, reports and audit records retain attributable state changes without treating routine clicks as academic competence.",
    ]:
        add_number(doc, step)

    add_heading(doc, "6. State Models", 1)
    add_table(
        doc,
        ["Record", "States", "Key transition controls"],
        [
            ("Practical resource", "Draft → Approved → Withdrawn", "Only Administrators approve or withdraw; withdrawal blocks future selection but not historical reading."),
            ("Task", "Draft → Published → Closed → Archived", "Only the assigned Lecturer publishes; major changes after publication create an auditable version."),
            ("Practical session", "Initiated → Active → Completed / Failed / Abandoned", "Session belongs to one Student, task and resource version; timings and agreed metadata are retained."),
            ("Submission", "Saved → Submitted → Returned → Resubmitted → Assessed", "Submission confirmation freezes a revision; resubmission creates a linked revision rather than overwriting evidence."),
            ("Result", "Draft assessment → Released", "Only the assigned Lecturer releases; the Student sees feedback and grade only after release."),
        ],
        widths=[1.3, 2.25, 3.55],
    )

    add_heading(doc, "7. Functional Requirements", 1)
    requirements = [
        ("FR-01", "Welcome and help", "The public interface shall present the platform identity, General Help, registration, sign-in and password recovery."),
        ("FR-02", "Registration", "The system shall accept only Student or Lecturer public registration and create a Pending Approval record."),
        ("FR-03", "Approval", "The Administrator shall approve, reject or suspend accounts with a recorded reason and audit event."),
        ("FR-04", "Academic setup", "The Administrator shall manage the academic hierarchy, course offerings, teaching assignments and enrolments."),
        ("FR-05", "Role workspaces", "Each authenticated user shall enter only the workspace and records permitted by role and academic relationship."),
        ("FR-06", "Resources", "The Administrator shall register providers, versions, launch methods, assumptions, status and permitted return fields for approved resources."),
        ("FR-07", "Tasks", "An assigned Lecturer shall create, publish, close and archive practical tasks with questions, deadlines, rubrics, evidence rules and optional approved-resource links."),
        ("FR-08", "Sessions", "The system shall retain Student, Lecturer, course, task, resource/version, run ID, timing, status, permitted parameters and agreed output metadata."),
        ("FR-09", "Submission", "A Student shall save work, upload permitted evidence, confirm submission and create a traceable resubmission when authorised."),
        ("FR-10", "Assessment", "The assigned Lecturer shall review the complete work trail, return feedback, grade, release and export results."),
        ("FR-11", "Notifications", "The system shall create in-platform notifications for significant academic state changes."),
        ("FR-12", "Reports", "Administrators shall filter and export management records; Lecturers shall export records only for assigned offerings."),
        ("FR-13", "Accounts", "Users shall manage permitted profile fields and use controlled password recovery; Administrators shall support documented resets."),
        ("FR-14", "Audit", "The system shall record actor, action, target, timestamp, reason where applicable and before/after state for sensitive actions."),
    ]
    add_table(doc, ["ID", "Area", "Requirement"], requirements, widths=[0.7, 1.25, 5.0], size=8.2)

    add_heading(doc, "8. Data Model", 1)
    data_rows = [
        ("Identity", "User, Role, AccountApproval, PasswordResetRecord, SessionRecord", "User 1—many approvals/events; role assigned once active"),
        ("Academic", "Faculty, Department, Programme, Course, AcademicPeriod, CourseOffering, TeachingAssignment, Enrolment", "CourseOffering joins course and period; assignments/enrolments grant scoped access"),
        ("Resources", "ResourceProvider, PracticalResource, ResourceVersion, ResourceApproval", "Only an approved version may be linked to a new task"),
        ("Assessment setup", "PracticalTask, TaskVersion, Rubric, RubricCriterion, EvidenceRequirement", "Task belongs to one offering and Lecturer; published changes are versioned"),
        ("Student activity", "PracticalSession, SessionParameter, SessionOutput, EvidenceFile, Submission, SubmissionRevision", "Every record traces to Student, task and applicable resource version"),
        ("Academic decision", "Assessment, CriterionScore, Feedback, ResultRelease", "Assessment targets one submission revision; release controls Student visibility"),
        ("Oversight", "Notification, AuditEvent, ExportRecord", "Sensitive changes remain attributable and time-stamped"),
    ]
    add_table(doc, ["Group", "Principal entities", "Relationship rule"], data_rows, widths=[1.05, 3.15, 2.8], size=8.0)
    add_para(doc, "Integrity rules: historical tasks, submissions, assessments and resource use are not hard-deleted through normal interfaces. Withdrawal or suspension changes current availability while preserving authorised historical access. Timestamps are stored in UTC and displayed in the configured institutional timezone.")

    add_heading(doc, "9. Approved-Resource Integration Contract", 1)
    add_table(
        doc,
        ["Direction", "Minimum fields", "Rule"],
        [
            ("Platform → resource", "launch_id, user pseudonymous identifier, role, course_offering_id, task_id, resource_version_id, expiry, permitted fields", "Short-lived signed launch; disclose only the context needed for the task."),
            ("Resource → platform", "launch_id, run_id, resource_version, start/end time, status, selected parameters, declared outputs/files, integrity value", "Accept only fields registered for that resource version; reject replay, expiry and mismatched ownership."),
            ("Controlled-link fallback", "task_id, resource_version, launch timestamp", "Where no data exchange exists, the Student submits evidence against the task and the platform does not infer completion from a click."),
        ],
        widths=[1.25, 3.65, 2.1],
        size=8.0,
    )

    add_heading(doc, "10. Page Inventory", 1)
    page_rows = [
        ("Public", "Welcome; General Help; Student/Lecturer registration; Sign in; Password recovery"),
        ("Approval/setup", "Pending status; Administrator approval queue; Academic setup form"),
        ("Administrator", "Dashboard; Users; Academic hierarchy; Course access; Resource catalogue; Records; Notifications; Reports; Password support; Audit log; Settings"),
        ("Lecturer", "Dashboard; Tasks; Task editor; Resource selection; Submission review; Grading and feedback; Result release; Export; Notifications; Profile"),
        ("Student", "Dashboard; Assigned tasks; Task detail; Virtual Modelling and Demonstration page; Session/evidence workspace; Submission confirmation; Revision history; Results; Notifications; Profile"),
    ]
    add_table(doc, ["Area", "Required pages"], page_rows, widths=[1.3, 5.7])

    add_heading(doc, "11. Security, Privacy and Accessibility Baseline", 1)
    controls = [
        "Django session authentication with HttpOnly, Secure and SameSite cookies; CSRF protection on every state-changing browser request.",
        "Argon2 preferred password hashing, password-strength checks, rate-limited sign-in and recovery, single-use expiring reset tokens and first-sign-in change for bootstrap credentials.",
        "Server-side role, course membership and object-ownership checks; denial responses must not disclose protected record existence.",
        "Allowlisted evidence types and size limits, generated storage names, content inspection where available, protected downloads and integrity metadata.",
        "No secrets, passwords, reset tokens or unnecessary personal content in source control or logs.",
        "Five-minute protected-session inactivity timeout, subject to formal approval in Section 14; unsaved forms should warn before expiry where feasible.",
        "WCAG 2.2 Level AA design target: keyboard operation, visible focus, meaningful labels, error identification, contrast, reflow, heading order and status announcements.",
        "Data minimisation, retention rules, documented backup and restoration, and test data that contains no real Student evidence.",
    ]
    for item in controls:
        add_bullet(doc, item)

    add_heading(doc, "12. Objective-to-Acceptance Traceability", 1)
    trace_rows = [
        ("Objective 1", "FR-04, FR-06, FR-07, FR-10, FR-12", "After Administrator setup, a Lecturer creates, publishes, reviews, grades, releases and exports one task for an assigned offering."),
        ("Objective 2", "FR-05, FR-08, FR-09, FR-11", "An enrolled Student opens the task, uses the approved resource, saves evidence, submits and traces status without seeing another Student’s work."),
        ("Objective 3", "FR-03–FR-14", "A complete task-to-result record is retrieved with valid relationships, versions, evidence, feedback and audit events."),
        ("Objective 4", "FR-06, FR-08", "Only an approved resource version launches from an assigned task; agreed return metadata is validated and linked to the correct session."),
    ]
    add_table(doc, ["Objective", "Requirements", "Minimum end-to-end acceptance"], trace_rows, widths=[1.0, 1.8, 4.2], size=8.2)

    add_heading(doc, "13. Development Increments and Definition of Done", 1)
    increments = [
        ("0", "Repository and environment", "Django project, settings separation, PostgreSQL, Docker Compose, CI, formatting and test harness."),
        ("1", "Identity and approval", "Registration → Pending Approval → academic setup → role dashboard; session and audit controls included."),
        ("2", "Academic structure", "Hierarchy, academic periods, offerings, teaching assignments and enrolments with scoped permissions."),
        ("3", "Resources and tasks", "Approved-resource catalogue and Lecturer draft/publish workflow."),
        ("4", "Sessions and submissions", "Virtual Modelling and Demonstration page, practical trail, protected evidence and revision-preserving submission."),
        ("5", "Assessment and results", "Lecturer review, feedback, grading, release and Student result visibility."),
        ("6", "Oversight", "Notifications, exports, reports, password support and audit review."),
        ("7", "Hardening and evaluation", "Security, accessibility, compatibility, performance and usability testing; Chapter Four evidence package."),
    ]
    add_table(doc, ["Increment", "Scope", "Exit condition"], increments, widths=[0.75, 1.65, 4.6], size=8.2)
    add_para(doc, "Definition of done for every increment: approved requirement traced; migration committed; server-side permissions tested; validation and error states implemented; audit event added where required; responsive and keyboard checks completed; automated tests passing; no real personal data in fixtures; documentation updated; and evidence retained for Chapter Four.")

    add_heading(doc, "14. Design-Freeze Decisions", 1)
    decisions = [
        ("Technology stack", "Approve the Django 5.2 LTS modular monolith and PostgreSQL 18 baseline.", "Pending approval"),
        ("Session timeout", "Retain the proposed five-minute inactivity timeout for protected academic sessions.", "Pending approval"),
        ("Assessment scale", "Store a configurable maximum score per task; do not hard-code one university-wide scale.", "Pending approval"),
        ("Evidence policy", "Initial allowlist: PDF, PNG, JPEG, CSV and TXT; proposed maximum 20 MB per file and 100 MB per submission.", "Pending approval"),
        ("Notifications", "Implement in-platform notifications first; email delivery remains an optional deployment feature.", "Pending approval"),
        ("Hosting", "Use localhost/Docker for development and evaluation; institutional HTTPS hosting is a later deployment decision.", "Pending approval"),
        ("Resource exchange", "Support controlled links first and a signed JSON return contract only for resources that can provide it.", "Pending approval"),
    ]
    add_table(doc, ["Decision", "Proposed baseline", "Status"], decisions, widths=[1.25, 4.55, 1.2], size=8.2)
    approval = doc.add_paragraph()
    approval.alignment = WD_ALIGN_PARAGRAPH.LEFT
    approval.add_run(
        "Approval record: Supervisor/authorised reviewer ____________________    "
        "Signature ____________________    Date ____________________"
    )

    add_heading(doc, "15. References", 1)
    refs = [
        "Django REST Framework. (2025). Version 3.16 announcement. https://www.django-rest-framework.org/community/3.16-announcement/",
        "Django Software Foundation. (2025). Django 5.2 release notes. https://docs.djangoproject.com/en/5.2/releases/5.2/",
        "Docker. (2026). Docker Compose. https://docs.docker.com/compose/",
        "OWASP Foundation. (2025). OWASP Application Security Verification Standard 5.0.0. https://owasp.org/www-project-application-security-verification-standard/",
        "PostgreSQL Global Development Group. (2026). PostgreSQL versioning policy. https://www.postgresql.org/support/versioning/",
        "World Wide Web Consortium. (2024). Web Content Accessibility Guidelines (WCAG) 2.2. https://www.w3.org/TR/WCAG22/",
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.line_spacing = 1.0

    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_SPEC)


def main() -> None:
    correct_chapters_one_two()
    rewrite_chapter_three()
    build_specification()
    print(OUT_CH12)
    print(OUT_CH3)
    print(OUT_SPEC)


if __name__ == "__main__":
    main()
